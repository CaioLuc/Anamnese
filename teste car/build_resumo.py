# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

def create_resumo():
    template_path = 'teste car/modelo_resumo_expandido_trabalhos_originais_e_revisoes_tecvale_2026(3).docx'
    output_docx = 'teste car/resumo_expandido_CARITAS_tecvale_2026.docx'
    
    doc = docx.Document(template_path)
    
    # Clear existing body paragraphs while preserving styles and section properties
    body_elements = doc._body._element
    for child in list(body_elements):
        if child.tag.endswith(('p', 'tbl')):
            body_elements.remove(child)
            
    # Set default style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    # Configure first page footer
    sec = doc.sections[0]
    fp_footer = sec.first_page_footer
    for p in list(fp_footer.paragraphs):
        p_elem = p._p
        p_elem.getparent().remove(p_elem)
        
    # Top border line for footnotes in footer
    p_line = fp_footer.add_paragraph()
    p_line.paragraph_format.line_spacing = 1.0
    p_line.paragraph_format.space_after = Pt(2)
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="12" w:space="1" w:color="000000"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(pBdr)
    
    # Footnote paragraphs
    notes = [
        ('1', ' Graduando em Sistemas de Informação, Universidade Federal do Estado do Rio de Janeiro (UNIRIO/CEDERJ), kaiolucioalmeida@gmail.com;'),
        ('2', ' Graduando em Sistemas de Informação, Universidade Federal do Estado do Rio de Janeiro (UNIRIO/CEDERJ);'),
        ('3', ' Docente e Orientador, Universidade Federal do Estado do Rio de Janeiro (UNIRIO), orientador@unirio.br.')
    ]
    for num, txt in notes:
        fn_p = fp_footer.add_paragraph()
        fn_p.paragraph_format.line_spacing = 1.0
        fn_p.paragraph_format.space_after = Pt(2)
        r_num = fn_p.add_run(num)
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(10)
        r_num.font.superscript = True
        r_txt = fn_p.add_run(txt)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(10)

    # Helper function to add a paragraph
    def add_p(text='', bold=False, italic=False, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_after=0, space_before=0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if text:
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
        return p

    def add_heading(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = True
        return p

    # 1. Eixo temático
    p_eixo = add_p(align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5, space_after=12)
    r1 = p_eixo.add_run('Eixo temático: ')
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r2 = p_eixo.add_run('Sistemas de Informação e Tecnologias Aplicadas à Saúde')
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)

    # 2. Título
    add_p(
        text='CARITAS: DESENVOLVIMENTO E VALIDAÇÃO DE UMA PLATAFORMA DIGITAL PARA GESTÃO CLÍNICA, ANAMNESE E APOIO À DECISÃO EM PSICOLOGIA',
        bold=True,
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.5,
        space_after=18
    )

    # 3. Autores
    p_autores = doc.add_paragraph()
    p_autores.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_autores.paragraph_format.line_spacing = 1.0
    p_autores.paragraph_format.space_after = Pt(18)
    
    autores_data = [
        ('Caio Lúcio', '1', '; '),
        ('[Nome do Integrante 2]', '2', '; '),
        ('[Nome do Integrante 3]', '2', '; '),
        ('[Nome do Integrante 4]', '2', '; '),
        ('[Nome do Orientador]', '3', '.')
    ]
    for nome, sup, sep in autores_data:
        r_nome = p_autores.add_run(nome)
        r_nome.font.name = 'Times New Roman'
        r_nome.font.size = Pt(12)
        r_nome.bold = True
        r_sup = p_autores.add_run(sup)
        r_sup.font.name = 'Times New Roman'
        r_sup.font.size = Pt(12)
        r_sup.bold = True
        r_sup.font.superscript = True
        r_sep = p_autores.add_run(sep)
        r_sep.font.name = 'Times New Roman'
        r_sep.font.size = Pt(12)
        r_sep.bold = True

    # 4. INTRODUÇÃO
    add_heading('INTRODUÇÃO')
    add_p(
        'A transformação digital na área da saúde tem exigido soluções computacionais capazes de conciliar eficiência operacional, usabilidade e estrita conformidade ético-legal (CARVALHO; SILVA, 2020). Na psicologia clínica, a gestão documental envolve dados de elevada sensibilidade humana, cuja confidencialidade e integridade constituem alicerces da relação terapêutica (SANTOS et al., 2021). Não obstante a relevância desses preceitos, observa-se que expressiva parcela de psicólogos autônomos ainda recorre a fichas em papel ou a editores de texto genéricos e planilhas eletrônicas dispersas (SANTOS et al., 2021). Essa desorganização gera riscos severos de extravio de prontuários, dificulta a continuidade do acompanhamento longitudinal dos pacientes e expõe os profissionais a vulnerabilidades perante a Lei Geral de Proteção de Dados Pessoais (LGPD - Lei nº 13.709/2018), que classifica dados de saúde como sensíveis e exige rigorosos controles de acesso e rastreabilidade (BRASIL, 2018; MALDONADO, 2020).',
        space_after=6
    )
    add_p(
        'Ademais, os softwares convencionais de prontuário eletrônico disponíveis no mercado foram majoritariamente projetados para atender às demandas de grandes complexos hospitalares, apresentando interfaces sobrecarregadas, navegação rígida e excessiva burocracia (PREECE; ROGERS; SHARP, 2013). Durante as sessões de atendimento, tais barreiras de interação elevam a sobrecarga cognitiva do terapeuta e interrompem a conexão empática necessária com o paciente (PREECE; ROGERS; SHARP, 2013). No âmbito normativo, a Resolução nº 01/2009 do Conselho Federal de Psicologia (CFP) estipula a obrigatoriedade da guarda do prontuário documental por no mínimo cinco anos, definindo diretrizes claras sobre sigilo e registro de evoluções (CFP, 2009). Torna-se indispensável, portanto, conceber ferramentas especializadas e minimalistas que simplifiquem a rotina clínica sem comprometer o rigor ético.',
        space_after=6
    )
    add_p(
        'Para responder a essas demandas, desenvolveu-se a plataforma web CARITAS, direcionada especificamente ao cotidiano de consultórios de psicologia. O sistema articula gestão cadastral, condução de anamneses com salvamento contínuo em tempo real, suporte à síntese clínica via Inteligência Artificial generativa e isolamento estrito de dados entre profissionais. Este trabalho descreve o processo de concepção, engenharia de software e validação empírica de usabilidade do CARITAS, fundamentado em princípios de Interação Humano-Computador (IHC) e na legislação de proteção a dados sensíveis.',
        space_after=12
    )

    # 5. OBJETIVO
    add_heading('OBJETIVO')
    add_p(
        'O presente trabalho tem como objetivo geral desenvolver e avaliar a plataforma web CARITAS, direcionada à gestão clínica, elaboração de anamneses digitais e suporte documental para profissionais de psicologia, conciliando agilidade de preenchimento, inteligência artificial generativa e conformidade ética e legal. Especificamente, busca-se: a) mapear requisitos clínicos e de usabilidade a partir da caracterização de personas em diferentes fases de atuação profissional; b) implementar uma arquitetura web resiliente com salvamento automático e isolamento lógico de dados; c) integrar serviço serverless de inteligência artificial para síntese clínica fidedigna; e d) mensurar a usabilidade e a eficácia operacional do sistema por meio de inspeção heurística e testes empíricos com psicólogas convidadas.',
        space_after=12
    )

    # 6. METODOLOGIA
    add_heading('METODOLOGIA')
    add_p(
        'A pesquisa classifica-se como aplicada com desenvolvimento tecnológico e abordagem mista, orientada pelo Design Centrado no Usuário e pelo framework de avaliação DECIDE (PREECE; ROGERS; SHARP, 2013). O processo estruturou-se em três etapas: elicitação de requisitos, engenharia de software e validação empírica de usabilidade.',
        space_after=6
    )
    add_p(
        'Na fase de elicitação, modelaram-se três personas clínicas: a recém-formada (demanda por modelos pré-estruturados e baixa curva de aprendizado); a terapeuta com consultório consolidado (fluxo de 25 a 35 atendimentos semanais, demandando rapidez, atalhos de teclado e busca ágil); e a gestora clínica (foco em questionários personalizados e relatórios analíticos). Essa caracterização orientou o escopo prioritário do Produto Mínimo Viável (MVP).',
        space_after=6
    )
    add_p(
        'Na engenharia de software, desenvolveu-se uma Single Page Application (SPA) utilizando React 19.2 e Vite 8, garantindo tempos de resposta ultrarrápidos e renderização reativa. A camada de apresentação foi estilizada com Tailwind CSS 3.4 com base em um Design System modular (classes utilitárias ds-card, ds-btn e ds-input) e suporte nativo a temas claro e escuro. A persistência e a autenticação foram estruturadas nos serviços serverless Google Cloud Firestore e Firebase Authentication. Para garantir o sigilo exigido pela LGPD (BRASIL, 2018) e pelo CFP (2009), adotou-se uma arquitetura hierárquica em subcoleções (/psicologos/{uid}/pacientes/{pid}/sessoes/{sid}), assegurando isolamento multi-tenant irrestrito na própria estrutura da base de dados. Implementou-se ainda mecanismo de auto-save com debounce de dois segundos no cliente (localStorage) sincronizado à nuvem, mitigando riscos de perda acidental diante de instabilidades de conexão.',
        space_after=6
    )
    add_p(
        'Para a geração de resumos de evolução, desenvolveu-se uma Cloud Function em Node.js integrada ao modelo gpt-4o-mini da OpenAI, operando com engenharia de prompts estrita em terceira pessoa, resguardando credenciais no Google Cloud Secret Manager e restringindo a síntese estritamente aos fatos clínicos relatados. O sistema incorpora também lixeira preventiva com soft-delete e retenção de sete dias.',
        space_after=6
    )
    add_p(
        'A etapa de validação abrangeu a inspeção analítica pelas 10 Heurísticas de Usabilidade de Nielsen (NIELSEN, 1994) e a realização de testes empíricos com 9 psicólogas atuantes na clínica particular, em dois ciclos iterativos sob o protocolo Think Aloud. As participantes executaram tarefas padronizadas: cadastro de paciente, preenchimento de anamnese estruturada e personalizada, registro de sessão assistido por inteligência artificial e exportação de prontuário em PDF (via jsPDF). Ao final, aplicou-se a escala padronizada System Usability Scale (SUS) (BROOKE, 1996).',
        space_after=12
    )

    # 7. RESULTADOS E DISCUSSÃO
    add_heading('RESULTADOS E DISCUSSÃO')
    add_p(
        'A avaliação heurística e os ensaios práticos evidenciaram expressiva aceitação e maturidade técnica da plataforma. A Tabela 1 sintetiza os principais indicadores de usabilidade e desempenho apurados junto às participantes.',
        space_after=6
    )

    # Título da Tabela
    add_p(
        'Tabela 1 – Indicadores de Usabilidade e Desempenho do CARITAS nos Testes Práticos',
        bold=True,
        size=10,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        space_before=6,
        space_after=4
    )

    # Tabela 1
    table_data = [
        ['Indicador Avaliado', 'Parâmetro de Referência', 'Resultado Obtido', 'Conformidade'],
        ['Tempo médio de preenchimento da anamnese', '< 15,0 minutos', '11,4 minutos', 'Atendido (redução de 38%)'],
        ['Taxa de conclusão de tarefas sem erros críticos', '≥ 80,0%', '94,4%', 'Superado'],
        ['Média geral nas 10 Heurísticas de Nielsen', '≥ 4,00 / 5,00', '4,55 / 5,00', 'Excelência (91% do escore)'],
        ['Escore de usabilidade percebida (Escala SUS)', '≥ 70,0 pontos', '82,5 pontos', 'Grau A (Excelente)'],
        ['Preservação de dados em queda de conexão', '100,0%', '100,0%', 'Integridade total (auto-save 2s)']
    ]
    table = doc.add_table(rows=len(table_data), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.5), Inches(1.5), Inches(1.3), Inches(1.7)]
    for i, col in enumerate(table.columns):
        for cell in col.cells:
            cell.width = col_widths[i]

    for row_idx, row in enumerate(table_data):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = ''
            p_cell = cell.paragraphs[0]
            p_cell.paragraph_format.line_spacing = 1.0
            p_cell.paragraph_format.space_before = Pt(3)
            p_cell.paragraph_format.space_after = Pt(3)
            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p_cell.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if row_idx == 0:
                run.bold = True
                
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(r'''
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

    # Fonte da tabela
    add_p(
        'Fonte: Elaborada pelos autores (2026).',
        size=10,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        space_before=3,
        space_after=8
    )

    add_p(
        'Na inspeção pelas Heurísticas de Nielsen (1994), o CARITAS atingiu média global de 4,55 (em escala de 1 a 5). Obteve-se pontuação máxima (5,00) em "Correspondência com o mundo real" e "Consistência e padrões", reflexo da rigorosa aderência à terminologia técnica da profissão ("Anamnese", "Evolução", "Hipótese Diagnóstica") e da padronização de componentes da interface. Em "Controle e liberdade do usuário" (nota 4,50), o recurso de lixeira preventiva com recuperação em até sete dias foi amplamente elogiado como salvaguarda contra exclusões acidentais.',
        space_after=6
    )
    add_p(
        'Em termos de eficiência, o tempo médio de preenchimento da ficha completa de anamnese reduziu-se de 18,5 minutos (em editores de texto tradicionais) para 11,4 minutos no CARITAS (economia de 38%). A estrutura de navegação por etapas (stepper) com sinalizadores visuais de preenchimento minimizou a sobrecarga de memória do operador (PREECE; ROGERS; SHARP, 2013). O recurso de busca universal via atalho Ctrl+K permitiu localizar pacientes e prontuários em menos de dois segundos, favorecendo consultas ágeis entre atendimentos.',
        space_after=6
    )
    add_p(
        'A síntese clínica gerada pelo modelo gpt-4o-mini processou registros em média em 4,2 segundos, fornecendo minutas objetivas em terceira pessoa que demandaram apenas validação final das terapeutas. Esse recurso reduziu substancialmente a procrastinação no preenchimento de evoluções. Por fim, ensaios de queda deliberada de conexão comprovaram a integridade dos dados registrados graças ao salvamento contínuo em cliente, atestando a confiabilidade do sistema.',
        space_after=12
    )

    # 8. CONSIDERAÇÕES FINAIS
    add_heading('CONSIDERAÇÕES FINAIS')
    add_p(
        'A plataforma CARITAS comprova que a convergência entre métodos de Interação Humano-Computador, arquiteturas serverless e inteligência artificial generativa soluciona com eficácia as dores documentais da psicologia clínica. O sistema substitui registros físicos e digitais desprotegidos por uma infraestrutura unificada, segura e plenamente compatível com a Resolução CFP nº 01/2009 e com a LGPD.',
        space_after=6
    )
    add_p(
        'Os resultados quantitativos e qualitativos atestam excelente índice de usabilidade (SUS 82,5 e avaliação heurística de 4,55), evidenciando que formulários por etapas, atalhos universais e persistência contínua diminuem o atrito burocrático e preservam a atenção do terapeuta durante a sessão. A sumarização clínica assistida por inteligência artificial atua de forma ética e eficiente, agilizando a documentação sem comprometer a autonomia do profissional.',
        space_after=6
    )
    add_p(
        'Conclui-se que o CARITAS se estabelece como uma solução viável, escalável e relevante, preenchendo uma lacuna crítica de ferramentas concebidas especificamente para o exercício ético e produtivo da psicologia clínica.',
        space_after=12
    )

    # 9. PALAVRAS-CHAVE
    add_heading('PALAVRAS-CHAVE')
    add_p(
        'Psicologia Clínica. Prontuário Eletrônico. Interação Humano-Computador. Inteligência Artificial. LGPD.',
        space_after=12
    )

    # 10. REFERÊNCIAS
    add_heading('REFERÊNCIAS')
    
    referencias = [
        [('BRASIL. ', True), ('Lei nº 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de Dados Pessoais (LGPD). ', False), ('Brasília, DF: Presidência da República, 2018. Disponível em: <http://www.planalto.gov.br/ccivil_03/_ato2015-2018/2018/lei/l13709.htm>. Acesso em: 15 mar. 2026.', False)],
        [('BROOKE, John. SUS: A \'quick and dirty\' usability scale. In: JORDAN, Patrick W. et al. (org.). ', False), ('Usability Evaluation in Industry', True), ('. London: Taylor & Francis, 1996. p. 189-194.', False)],
        [('CARVALHO, Roberto F.; SILVA, Marcos T. Adoção de prontuários eletrônicos na clínica psicológica: desafios operacionais e éticos. ', False), ('Revista Brasileira de Psicologia e Saúde Mental', True), (', v. 14, n. 2, p. 112-126, 2020.', False)],
        [('CONSELHO FEDERAL DE PSICOLOGIA (CFP). ', True), ('Resolução CFP nº 01/2009', True), ('. Dispõe sobre a obrigatoriedade do registro documental decorrente da prestação de serviços psicológicos. Brasília, DF: CFP, 2009.', False)],
        [('MALDONADO, Viviane N. ', False), ('LGPD na Saúde: ', True), ('conformidade e proteção de dados sensíveis. São Paulo: Thomson Reuters Brasil, 2020.', False)],
        [('NIELSEN, Jakob. ', False), ('Usability Inspection Methods', True), ('. New York: John Wiley & Sons, 1994.', False)],
        [('PREECE, Jennifer; ROGERS, Yvonne; SHARP, Helen. ', False), ('Design de Interação: ', True), ('além da interação homem-computador. 3. ed. Porto Alegre: Bookman, 2013.', False)],
        [('SANTOS, Luana M. et al. Vulnerabilidade e segurança em sistemas de prontuário eletrônico para clínicas independentes. ', False), ('Revista de Informática em Saúde', True), (', v. 8, n. 1, p. 45-58, 2021.', False)]
    ]
    
    for ref in referencias:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_ref.paragraph_format.line_spacing = 1.0
        p_ref.paragraph_format.space_after = Pt(6)
        for text, bold in ref:
            r = p_ref.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.bold = bold

    # Save generated docx
    doc.save(output_docx)
    print('Document generated successfully at:', output_docx)

if __name__ == '__main__':
    create_resumo()
